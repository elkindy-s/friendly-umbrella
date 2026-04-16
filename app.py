from flask import Flask, render_template, request, redirect, url_for, send_file, session
import os
import pandas as pd
import re
import easyocr
from datetime import datetime
from docx import Document
from docx.shared import Pt
from openpyxl import load_workbook
from openpyxl.styles import Font

app = Flask(__name__)
app.secret_key = "yoursecretkey"

# Folders
UPLOAD_FOLDER = 'uploads'
OUTPUT_FOLDER = 'outputs'
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
os.makedirs(OUTPUT_FOLDER, exist_ok=True)

# Login credentials
USERNAME = "Revenue Starehe"
PASSWORD = "masky"

def extract_field(text, pattern):
    match = re.search(pattern, text, re.IGNORECASE | re.DOTALL)
    return match.group(1).strip() if match else ''

@app.route("/", methods=["GET"])
def landing():
    return render_template("landing.html")

@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        if request.form["username"] == USERNAME and request.form["password"] == PASSWORD:
            session["logged_in"] = True
            return redirect(url_for("upload"))
    return render_template("login.html")

@app.route("/upload", methods=["GET", "POST"])
def upload():
    if not session.get("logged_in"):
        return redirect(url_for("login"))

    if request.method == "POST":
        files = request.files.getlist("files[]")
        data = []
        reader = easyocr.Reader(['en'])

        for file in files:
            if file.filename:
                filepath = os.path.join(UPLOAD_FOLDER, file.filename)
                file.save(filepath)
                result = reader.readtext(filepath, detail=0)
                text = " ".join(result)

                parcel = extract_field(text, r'Parcel Number[:\s]*([\w/]+)')
                phones = re.findall(r'\b\d{9,12}\b', text)
                phone = phones[0] if phones else ''
                property_name = extract_field(text, r'Name[:\s]*([\w\s]+)')
                bill_no = extract_field(text, r'Bill No[:\s]*([\w\-]+)')
                total = extract_field(text, r'Bill Total Amount\s*([\d,]+(?:\.\d{2})?)')

                data.append({
                    'Parcel Number': parcel,
                    'Phone Number': phone,
                    'Property Name': property_name,
                    'Bill Number': bill_no,
                    'Bill Total': total
                })

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        excel_filename = os.path.join(OUTPUT_FOLDER, f"output_{timestamp}.xlsx")
        word_filename = os.path.join(OUTPUT_FOLDER, f"output_{timestamp}.docx")

        # Save Excel with Times New Roman
        df = pd.DataFrame(data)
        df.to_excel(excel_filename, index=False)
        wb = load_workbook(excel_filename)
        ws = wb.active
        for cell in ws[1]:
            cell.font = Font(name="Times New Roman", bold=True)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="Times New Roman")
        wb.save(excel_filename)

        # Save Word with Times New Roman
        doc = Document()
        style = doc.styles['Normal']
        font = style.font
        font.name = 'Times New Roman'
        font.size = Pt(12)
        table = doc.add_table(rows=1, cols=len(df.columns))
        hdr_cells = table.rows[0].cells
        for i, col in enumerate(df.columns):
            hdr_cells[i].text = col
        for row in df.itertuples(index=False):
            row_cells = table.add_row().cells
            for i, value in enumerate(row):
                row_cells[i].text = str(value)
        doc.save(word_filename)

        return render_template("upload.html", excel_file=excel_filename, word_file=word_filename)

    return render_template("upload.html")

@app.route("/download/<path:filename>")
def download(filename):
    return send_file(filename, as_attachment=True)

if __name__ == "__main__":
    app.run(debug=True)
